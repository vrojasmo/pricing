"""
Generates Excel + CSV + console reports from annotated ProductPrice records.
"""

from __future__ import annotations

import csv
import json
import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

import pandas as pd

from .scrapers.base import ProductPrice

logger = logging.getLogger(__name__)

RED   = "\033[91m"
GREEN = "\033[92m"
YELLOW = "\033[93m"
RESET = "\033[0m"


def _to_df(products: list[ProductPrice]) -> pd.DataFrame:
    rows = []
    for p in products:
        rows.append({
            "Retailer":        p.retailer,
            "Country":         p.country,
            "Product Name":    p.product_name,
            "Retailer Price (€)": p.price,
            "Official Price (€)": p.official_price,
            "Discount (%)":    p.discount_pct,
            "Compliant":       (
                "YES" if p.compliant is True
                else "NO" if p.compliant is False
                else "UNKNOWN"
            ),
            "URL":             p.url,
            "Scraped At":      p.scraped_at,
        })
    return pd.DataFrame(rows)


def save_excel(products: list[ProductPrice], output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    path = output_dir / f"bover_prices_{ts}.xlsx"

    df = _to_df(products)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        # ── Full data sheet ────────────────────────────────────────────────
        df.to_excel(writer, sheet_name="All Products", index=False)
        _format_excel_sheet(writer, "All Products", df)

        # ── Non-compliant only ─────────────────────────────────────────────
        non_compliant = df[df["Compliant"] == "NO"]
        if not non_compliant.empty:
            non_compliant.to_excel(writer, sheet_name="NON-COMPLIANT", index=False)
            _format_excel_sheet(writer, "NON-COMPLIANT", non_compliant)

        # ── Summary by retailer ────────────────────────────────────────────
        summary = _build_summary(df)
        summary.to_excel(writer, sheet_name="Summary", index=False)

        # ── By country sheets ──────────────────────────────────────────────
        for country in df["Country"].unique():
            cdf = df[df["Country"] == country]
            cdf.to_excel(writer, sheet_name=f"Country_{country}", index=False)
            _format_excel_sheet(writer, f"Country_{country}", cdf)

    logger.info("Excel report saved: %s", path)
    return path


def _format_excel_sheet(writer, sheet_name: str, df: pd.DataFrame) -> None:
    from openpyxl.styles import PatternFill, Font, Alignment
    from openpyxl.utils import get_column_letter

    ws = writer.sheets[sheet_name]

    # Header style
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Row colours
    red_fill   = PatternFill("solid", fgColor="FFCCCC")
    green_fill = PatternFill("solid", fgColor="CCFFCC")
    grey_fill  = PatternFill("solid", fgColor="F2F2F2")

    compliant_col = None
    for idx, cell in enumerate(ws[1], 1):
        if cell.value == "Compliant":
            compliant_col = idx
            break

    for row_idx, row in enumerate(ws.iter_rows(min_row=2), 2):
        if compliant_col:
            val = ws.cell(row=row_idx, column=compliant_col).value
            fill = red_fill if val == "NO" else green_fill if val == "YES" else grey_fill
            for cell in row:
                cell.fill = fill

    # Auto-width
    for col_idx, col in enumerate(ws.columns, 1):
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 60)


def _build_summary(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for retailer in df["Retailer"].unique():
        rdf = df[df["Retailer"] == retailer]
        country = rdf["Country"].iloc[0]
        total = len(rdf)
        matched = (rdf["Compliant"] != "UNKNOWN").sum()
        compliant = (rdf["Compliant"] == "YES").sum()
        non_compliant = (rdf["Compliant"] == "NO").sum()
        unknown = (rdf["Compliant"] == "UNKNOWN").sum()
        max_disc = rdf["Discount (%)"].max() if matched else None
        avg_disc = round(rdf[rdf["Compliant"] != "UNKNOWN"]["Discount (%)"].mean(), 2) if matched else None
        rows.append({
            "Retailer":           retailer,
            "Country":            country,
            "Total Products":     total,
            "Matched to Official": matched,
            "Compliant":          compliant,
            "NON-COMPLIANT":      non_compliant,
            "Unknown":            unknown,
            "Max Discount (%)":   max_disc,
            "Avg Discount (%)":   avg_disc,
            "Compliance Rate (%)": round(compliant / matched * 100, 1) if matched else None,
        })
    return pd.DataFrame(rows).sort_values("NON-COMPLIANT", ascending=False)


def save_csv(products: list[ProductPrice], output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    path = output_dir / f"bover_prices_{ts}.csv"
    df = _to_df(products)
    df.to_csv(path, index=False)
    logger.info("CSV saved: %s", path)
    return path


def save_json(products: list[ProductPrice], output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    path = output_dir / f"bover_prices_{ts}.json"
    data = [
        {
            "retailer": p.retailer,
            "country": p.country,
            "product_name": p.product_name,
            "price": p.price,
            "official_price": p.official_price,
            "discount_pct": p.discount_pct,
            "compliant": p.compliant,
            "url": p.url,
            "scraped_at": p.scraped_at,
        }
        for p in products
    ]
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False))
    logger.info("JSON saved: %s", path)
    return path


def print_console_summary(products: list[ProductPrice]) -> None:
    """Print a colour-coded summary table to stdout."""
    if not products:
        print("No products scraped.")
        return

    total      = len(products)
    matched    = sum(1 for p in products if p.official_price is not None)
    compliant  = sum(1 for p in products if p.compliant is True)
    non_comp   = sum(1 for p in products if p.compliant is False)
    unknown    = total - matched

    print("\n" + "=" * 70)
    print("  BOVER PRICE COMPLIANCE REPORT")
    print("=" * 70)
    print(f"  Total products scraped : {total}")
    print(f"  Matched to official PL : {matched}")
    print(f"  {GREEN}Compliant              : {compliant}{RESET}")
    print(f"  {RED}NON-COMPLIANT          : {non_comp}{RESET}")
    print(f"  {YELLOW}Unknown (not in PL)    : {unknown}{RESET}")
    print("=" * 70)

    # Retailers summary
    retailers: dict[str, dict] = {}
    for p in products:
        r = retailers.setdefault(p.retailer, {
            "country": p.country,
            "total": 0, "compliant": 0, "non_compliant": 0, "unknown": 0,
        })
        r["total"] += 1
        if p.compliant is True:
            r["compliant"] += 1
        elif p.compliant is False:
            r["non_compliant"] += 1
        else:
            r["unknown"] += 1

    print(f"\n{'RETAILER':<28} {'CTR':<5} {'TOTAL':>6} {'OK':>5} {'FAIL':>5} {'UNK':>5}")
    print("-" * 55)
    for name, s in sorted(retailers.items(), key=lambda x: -x[1]["non_compliant"]):
        colour = RED if s["non_compliant"] > 0 else GREEN
        print(
            f"{colour}{name:<28}{RESET} {s['country']:<5} "
            f"{s['total']:>6} {s['compliant']:>5} {s['non_compliant']:>5} {s['unknown']:>5}"
        )

    # Non-compliant detail
    violations = [p for p in products if p.compliant is False]
    if violations:
        print(f"\n{RED}── NON-COMPLIANT PRODUCTS ──────────────────────────────────────{RESET}")
        for p in sorted(violations, key=lambda x: x.discount_pct or 0, reverse=True):
            print(
                f"  [{p.country}] {p.retailer:<28} | "
                f"{p.product_name[:40]:<40} | "
                f"€{p.price:>8.2f} vs official €{p.official_price:>8.2f} | "
                f"{RED}-{p.discount_pct:.1f}%{RESET}"
            )

    print("")
